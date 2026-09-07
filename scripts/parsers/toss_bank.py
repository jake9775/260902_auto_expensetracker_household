# -*- coding: utf-8 -*-
"""
토스뱅크 '거래내역' 엑셀 파일 파서.

기대하는 형식 (토스뱅크 앱/홈페이지에서 내려받은 그대로):
  1~4행: 제목/성명/계좌번호/조회기간 안내문
  5~8행: 안내 문구, 빈 줄
  9행  : 헤더 (거래 일시 적요 거래 유형 거래 기관 계좌번호 거래 금액 거래 후 잔액 메모)
  10행~: 실제 거래 데이터

  줄 수가 계좌 종류(개인 계좌/모임통장 등)에 따라 조금 달라질 수 있어서,
  고정된 줄 번호 대신 "거래 일시"라는 글자가 있는 헤더 행을 찾아서 그 다음
  줄부터 데이터로 읽는다.

  우리은행/현대카드와 달리 출금/입금이 한 칸(거래 금액)에 부호로 구분되어
  있다 (출금 = 음수, 입금 = 양수).

비밀번호:
  토스뱅크에서 받은 엑셀 파일은 보통 비밀번호가 걸려있다. 파일이 잠겨있으면
  실행 중 검은 창에서 비밀번호를 직접 입력받는다 (저장하지 않고 실행할
  때마다 새로 물어본다).

  화면에 안 보이게 입력받는 getpass 방식은 이 배치파일의 한글 설정
  (chcp 65001)과 겹치면 일부 윈도우 환경에서 키 입력을 아예 못 받는
  문제가 있어서, 그냥 화면에 보이는 일반 입력(input)으로 받는다.
"""
import datetime
import io

import msoffcrypto
import openpyxl


def _try_decrypt(path, password):
    with open(path, "rb") as f:
        office_file = msoffcrypto.OfficeFile(f)
        office_file.load_key(password=password)
        decrypted = io.BytesIO()
        office_file.decrypt(decrypted)
    return openpyxl.load_workbook(decrypted, data_only=True)


def _open_workbook(path):
    """비밀번호가 걸려있지 않으면 그냥 열고, 걸려있으면 사용자에게 물어봐서 연다.
    (저장은 하지 않고 실행할 때마다 새로 입력받는다)"""
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception:
        pass  # 비밀번호로 암호화된 파일은 openpyxl이 바로 열지 못한다.

    for _ in range(3):
        password = input(f"  '{path.name}' 파일은 비밀번호로 잠겨있습니다. 비밀번호를 입력하세요: ")
        try:
            return _try_decrypt(path, password)
        except Exception:
            print("  [알림] 비밀번호가 맞지 않는 것 같습니다. 다시 시도해주세요.")

    raise RuntimeError(f"'{path.name}' 파일 비밀번호를 확인하지 못해 열 수 없습니다.")


def parse(path):
    wb = _open_workbook(path)
    ws = wb[wb.sheetnames[0]]

    all_rows = list(ws.iter_rows(values_only=True))

    header_idx = None
    for i, raw in enumerate(all_rows):
        if raw and len(raw) > 1 and raw[1] == "거래 일시":
            header_idx = i
            break
    if header_idx is None:
        print(f"  [알림] '{path.name}' 에서 헤더 행('거래 일시')을 찾지 못했습니다. 건너뜁니다.")
        return []

    rows = []
    for raw in all_rows[header_idx + 1:]:
        if raw is None or len(raw) < 7:
            continue

        dt_str = raw[1]
        if not dt_str:
            continue
        try:
            dt = datetime.datetime.strptime(str(dt_str).strip(), "%Y.%m.%d %H:%M:%S")
        except ValueError:
            continue

        name = str(raw[2] or "").strip()      # 적요: 가맹점명 또는 상대방 이름
        tx_type = str(raw[3] or "").strip()    # 거래 유형: 체크카드결제/입금/모임원송금/이자입금 등
        amount = raw[6]
        if amount in (None, ""):
            continue
        amount = float(amount)

        # 카드결제는 적요만으로 충분히 알아볼 수 있지만, 계좌이체/입금류는
        # 상대방 이름만으로는 헷갈릴 수 있어 거래 유형을 괄호로 덧붙인다.
        if tx_type and tx_type != "체크카드결제":
            desc = f"{name}({tx_type})" if name else tx_type
        else:
            desc = name

        row = {
            "은행": "토스뱅크",
            "거래일자": dt.date(),
            "거래시간": dt.time(),
            "적요": desc,
            "내용": desc,
        }
        if amount < 0:
            row["출금"], row["입금"] = int(round(-amount)), 0
        else:
            row["출금"], row["입금"] = 0, int(round(amount))

        memo = str(raw[8] or "").strip() if len(raw) > 8 else ""
        if memo:
            row["메모"] = memo  # 토스뱅크 파일에 이미 적혀있던 메모를 그대로 살려서 씀

        rows.append(row)

    return rows
