# -*- coding: utf-8 -*-
"""
구글플레이 결제내역 엑셀 파일 파서.

naver_pay.py / coupang.py 와 마찬가지로 별도 거래 행을 만들지 않고,
(날짜, 시각, 금액) -> 상품명 조회용 목록만 만든다.

헤더: 결제일시, PG사, 상점명, 상품명, 결제금액, 구분, 주문번호, 결제방법
"""
import datetime
import openpyxl


def extract_entries(path):
    """[(date, hour, minute, amount, product_name), ...] 목록을 돌려준다.
    '결제'가 아닌 항목(환불/취소 등)은 제외한다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    entries = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or len(raw) < 6:
            continue
        pay_dt, _pg, _shop, product, amount, kind = raw[:6]
        if not pay_dt or amount in (None, "") or kind != "결제":
            continue
        dt = _parse_datetime(pay_dt)
        if dt is None:
            continue
        product = str(product or "").strip()
        entries.append((dt.date(), dt.hour, dt.minute, int(amount), product))
    return entries


def _parse_datetime(value):
    if isinstance(value, datetime.datetime):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return datetime.datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    return None
