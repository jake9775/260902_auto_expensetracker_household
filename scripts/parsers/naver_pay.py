# -*- coding: utf-8 -*-
"""
네이버페이 결제내역 엑셀 파일 파서.

이 파일은 별도의 계좌/카드가 아니라, 우리은행·현대카드 내역에 "네이버페이"로
뭉뚱그려 찍히는 결제 건의 실제 상품명을 알려주는 "보조 자료"로만 쓴다.
그래서 거래 행을 만들지 않고, (날짜, 시각, 금액) -> 상품명 조회용 목록만 만든다.

헤더: 승인번호, 카드사, 카드번호(유효기간), 거래종류/할부, 결제일자, 취소일자,
      상품명, 승인금액, 취소금액, 공급가액, 부가세액, 봉사료, 컵보증금, 합계
"""
import datetime
import openpyxl


def extract_entries(path):
    """[(date, hour, minute, amount, product_name), ...] 목록을 돌려준다.

    쿠팡처럼, 한 결제(승인번호)에 서로 다른 판매자 상품이 섞여 있으면 네이버페이도
    같은 승인번호로 여러 줄에 나눠서 찍힐 수 있다(거래일시·승인금액은 줄마다 동일).
    그대로 두면 (날짜, 금액) 키가 우연히 같은 것처럼 겹쳐서 상품명 하나만 뽑히고
    나머지가 유실되므로, 같은 승인번호끼리는 상품명을 합쳐서 하나의 항목으로 만든다.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    by_approval = {}  # 승인번호 -> 항목 dict
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or len(raw) < 8:
            continue
        approval_no, _card_co, _card_no, _txn_type, pay_dt, _cancel_dt, product, amount = raw[:8]
        if not approval_no or not pay_dt or not amount:
            continue
        dt = _parse_datetime(pay_dt)
        if dt is None:
            continue
        product = str(product or "").strip()

        if approval_no in by_approval:
            by_approval[approval_no]["products"].append(product)
        else:
            by_approval[approval_no] = {
                "date": dt.date(), "hour": dt.hour, "minute": dt.minute,
                "amount": int(amount), "products": [product],
            }

    entries = []
    for entry in by_approval.values():
        product = " + ".join(dict.fromkeys(entry["products"]))  # 중복 제거하며 합침
        entries.append((entry["date"], entry["hour"], entry["minute"], entry["amount"], product))
    return entries


def _parse_datetime(value):
    if isinstance(value, datetime.datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.datetime.strptime(value.strip(), "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None
    return None
