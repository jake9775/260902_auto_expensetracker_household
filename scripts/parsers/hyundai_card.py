# -*- coding: utf-8 -*-
"""
현대카드 '실시간 이용내역' 엑셀 파일 파서.

기대하는 형식:
  1~2행: 제목
  3행  : 헤더 (승인일 승인시각 카드구분 카드종류 가맹점명 승인금액 이용구분 할부개월 승인번호 취소일 승인구분)
  4행~ : 실제 승인 데이터, 맨 아래에 '국내 일시불 소계' 등 요약행이 붙어있음

주의: '승인일'은 엑셀 표시형식이 날짜가 아니라 일반 숫자(엑셀 시리얼 넘버)로
저장되어 있어서, 그냥 읽으면 46262 같은 숫자로 나온다. 엑셀 기준일(1899-12-30)
기준으로 직접 날짜로 변환해줘야 한다.
"""
import datetime
import openpyxl

_EXCEL_EPOCH = datetime.date(1899, 12, 30)


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    for raw in ws.iter_rows(min_row=4, values_only=True):
        if raw is None or len(raw) < 6:
            continue
        approve_date, approve_time, _card_gubun, _card_type, merchant, amount = raw[:6]
        approve_type = raw[10] if len(raw) > 10 else None

        if not isinstance(approve_date, (int, float)):
            continue  # '국내 일시불 소계' 등 요약행 ('-' 표시) 은 건너뜀
        if not merchant or amount in (None, ""):
            continue

        d = _EXCEL_EPOCH + datetime.timedelta(days=int(approve_date))
        t = approve_time if isinstance(approve_time, datetime.time) else datetime.time(0, 0)
        desc = str(merchant).strip()

        row = {
            "은행": "현대카드",
            "거래일자": d,
            "거래시간": t,
            "적요": desc,
            "내용": desc,
        }

        is_cancel = isinstance(approve_type, str) and "취소" in approve_type
        if is_cancel:
            row["출금"], row["입금"] = 0, amount
            row["대분류"], row["소분류"] = "환불", "카드취소"
        else:
            row["출금"], row["입금"] = amount, 0

        rows.append(row)

    return rows
