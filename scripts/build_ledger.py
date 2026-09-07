# -*- coding: utf-8 -*-
"""
가계부 반자동화 메인 스크립트.

사용법: `input/` 폴더에 이번 달 은행/카드/네이버페이/쿠팡 파일을 넣고
        `가계부_업데이트.bat`을 더블클릭 (또는 `python scripts/build_ledger.py`).

동작:
  1. input/ 폴더의 파일을 이름으로 인식해서 종류별 파서로 읽는다.
  2. 우리은행/현대카드 내역을 공통 형식으로 정리한다.
  3. 네이버페이/쿠팡 내역은 "상품명 보충용" 자료로만 써서 메모 칸을 채운다.
  4. (AUTO_FILL=True일 때만) rules/categories.csv 규칙으로 대분류/소분류를 자동으로 채운다.
  5. output/가계부.xlsx 에 기존 내용과 합쳐서(중복 제외) 새로 저장한다.
"""
import csv
import datetime
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parsers import coupang, google_play, hyundai_card, naver_pay, toss_bank, woori_bank  # noqa: E402

BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_DIR = BASE_DIR / "input"
RULES_PATH = BASE_DIR / "rules" / "categories.csv"
MEMO_RULES_PATH = BASE_DIR / "rules" / "memo_mapping.csv"
OUTPUT_PATH = BASE_DIR / "output" / "가계부.xlsx"

COLUMNS = ["은행", "거래일자", "거래시간", "적요", "출금", "입금", "내용", "대분류", "소분류", "메모"]

# False로 두면 '내용/대분류/소분류'는 자동으로 채우지 않고 전부 빈칸으로 저장한다.
# (지금은 사용자가 직접 채우기로 해서 꺼둠. 나중에 자동 분류를 다시 쓰고 싶으면 True로 바꾸면 됨)
AUTO_FILL = False

# 파일명에 이 키워드가 들어있으면 해당 파서로 처리한다. (엑셀 거래 내역 파서)
# (같은 은행/카드사라도 실제 다운로드 파일명이 한글/영문 등 여러 형태로 나올 수 있어서
#  키워드를 여러 개 등록해둔다. 대소문자는 구분하지 않는다.)
LEDGER_FILE_RULES = [
    ("우리은행", woori_bank.parse),
    ("거래내역조회", woori_bank.parse),  # 우리은행 인터넷뱅킹 기본 다운로드 파일명
    ("현대카드", hyundai_card.parse),
    ("hyundaicard", hyundai_card.parse),
    ("토스뱅크", toss_bank.parse),
]
# 파일명에 이 키워드가 들어있으면 "상품명 보충용" 조회 자료로만 쓴다.
LOOKUP_FILE_RULES = [
    ("네이버페이", naver_pay.extract_entries),
    ("npay", naver_pay.extract_entries),
    ("쿠팡", coupang.extract_entries),
    ("구글플레이", google_play.extract_entries),
]
# 아직 파서가 없는 은행/카드사 - 실제 파일을 받으면 parsers/ 에 파서를 추가하고
# 여기 LEDGER_FILE_RULES 에 한 줄만 추가하면 됨.
KNOWN_BUT_UNSUPPORTED = ["모임통장", "모임체크카드"]


def discover_and_load():
    """input/ 폴더를 훑어서 (거래행 목록, 상품명 보충용 목록) 을 돌려준다."""
    ledger_rows = []
    lookup_entries = []  # [(date, hour, minute, amount, product), ...]

    if not INPUT_DIR.exists():
        print(f"[경고] {INPUT_DIR} 폴더가 없습니다.")
        return ledger_rows, lookup_entries

    files = [
        p for p in INPUT_DIR.iterdir()
        if p.is_file() and not p.name.startswith("~$") and not p.name.startswith(".")
    ]
    if not files:
        print(f"[안내] {INPUT_DIR} 폴더에 파일이 없습니다. 이번 달 파일을 넣고 다시 실행해주세요.")

    for path in files:
        name = path.name
        matched = False

        name_lower = name.lower()

        for keyword, parse_fn in LEDGER_FILE_RULES:
            if keyword.lower() in name_lower:
                print(f"  - 거래내역으로 인식: {name} ({keyword})")
                ledger_rows.extend(parse_fn(path))
                matched = True
                break
        if matched:
            continue

        for keyword, extract_fn in LOOKUP_FILE_RULES:
            if keyword.lower() in name_lower:
                print(f"  - 상품명 보충 자료로 인식: {name} ({keyword})")
                lookup_entries.extend(extract_fn(path))
                matched = True
                break
        if matched:
            continue

        for keyword in KNOWN_BUT_UNSUPPORTED:
            if keyword in name:
                print(f"  [알림] '{name}' 은(는) 아직 지원하지 않는 형식입니다. "
                      f"scripts/parsers 에 {keyword} 파서를 추가해주세요. 이번엔 건너뜁니다.")
                matched = True
                break
        if matched:
            continue

        print(f"  [알림] '{name}' 파일 종류를 알 수 없어 건너뜁니다. "
              f"(파일명에 은행/카드사 이름이 들어있는지 확인해주세요)")

    return ledger_rows, lookup_entries


def build_amount_lookup(entries):
    """(날짜, 금액) -> [(분단위 시각, 상품명), ...] 형태로 묶는다."""
    lookup = {}
    for date, hour, minute, amount, product in entries:
        if not product:
            continue
        lookup.setdefault((date, amount), []).append((hour * 60 + minute, product))
    return lookup


_MEMO_KEYWORDS = ("쿠팡", "네이버페이", "네이버 페이", "네이버", "구글플레이")


def enrich_memo(rows, amount_lookup):
    """쿠팡/네이버페이/구글플레이로 결제된 항목은 실제 상품명을 '메모' 칸에 적어준다.
    (날짜/시간/금액이 일치하는 상세 내역을 찾아서 채움. 못 찾으면 빈칸으로 둠)

    카드사가 표시하는 '승인일'과 실제 결제 영수증(쿠팡/네이버페이/구글플레이)의 날짜가
    하루 정도 어긋나는 경우가 있어서(예: 자정 근처 결제, 카드사 정산 기준일 차이),
    같은 날 뿐 아니라 하루 전/하루 후까지 넓혀서 찾는다. 그중 시각이 가장 가까운
    걸 고른다.

    토스뱅크 파일 자체에 적혀있던 메모처럼 이미 메모가 채워진 행은 건드리지
    않는다. rules/memo_mapping.csv 고정 매핑(apply_memo_rules)보다 먼저
    실행해서, 실제 주문내역으로 확인된 상품명을 고정 매핑보다 우선 쓴다.

    현대카드로 결제된 항목은 전부 앱/서비스 정기결제(구독)라서, 자동으로 찾은
    상품명 뒤에 '구독'을 붙여준다. (이미 '구독'으로 끝나면 중복으로 붙이지 않음)

    이미 가계부.xlsx에 저장되어 있던 과거 거래(existing_rows)에도 이 함수를
    다시 돌릴 수 있는데, 그 경우 거래일자/거래시간 칸에 날짜(date)/시각(time)이
    아니라 둘 다 같은 datetime 값이 들어있다 (엑셀에 저장할 때 두 칸에 같은
    datetime을 넣고 표시 형식만 다르게 줬기 때문). 그대로 두면 날짜 비교가
    안 맞아서 매칭이 하나도 안 되므로, 먼저 date/time으로 변환해서 쓴다."""
    for row in rows:
        row.setdefault("메모", "")
        if row["메모"]:
            continue
        if not any(kw in row["적요"] for kw in _MEMO_KEYWORDS):
            continue
        amount = row["출금"] or row["입금"]
        tx_date = row["거래일자"]
        if isinstance(tx_date, datetime.datetime):
            tx_date = tx_date.date()
        tx_time = row["거래시간"]
        if isinstance(tx_time, datetime.datetime):
            tx_time = tx_time.time()
        candidates = []
        for delta in (0, -1, 1):
            day = tx_date + datetime.timedelta(days=delta)
            candidates.extend(amount_lookup.get((day, amount), []))
        if not candidates:
            continue
        row_minutes = tx_time.hour * 60 + tx_time.minute
        _, best_product = min(candidates, key=lambda c: abs(c[0] - row_minutes))
        if row.get("은행") == "현대카드" and not best_product.endswith("구독"):
            best_product = f"{best_product} 구독"
        row["메모"] = best_product


def load_memo_rules():
    """rules/memo_mapping.csv 를 읽는다. (가맹점명 키워드 -> 메모 고정 매핑)
    구글플레이처럼 매달 내역이 바뀌는 게 아니라, 항상 같은 뜻인 가맹점명
    (예: 'ktIOT자동이체' = 차량 커넥티드 서비스 요금)을 고정으로 채울 때 쓴다."""
    if not MEMO_RULES_PATH.exists():
        return []
    rules = []
    with open(MEMO_RULES_PATH, encoding="utf-8-sig") as f:
        for item in csv.DictReader(f):
            keyword = (item.get("키워드") or "").strip()
            memo = (item.get("메모") or "").strip()
            if keyword and memo:
                rules.append((keyword, memo))
    return rules


def apply_memo_rules(rows, memo_rules):
    """rules/memo_mapping.csv 고정 매핑을 적용한다. enrich_memo(쿠팡/네이버페이/
    구글플레이 실제 주문내역으로 찾은 상품명)보다 나중에 실행해서, 실제 주문내역이
    있으면 그걸 우선 쓰고, 그래도 메모가 비어있는 행에만 고정 매핑을 채운다.

    토스뱅크 파서처럼 파일 자체에 이미 사용자가 적어둔 메모가 있는 행도
    (가장 정확한 정보이므로) 고정 매핑으로 덮어쓰지 않고 그대로 둔다."""
    for row in rows:
        row.setdefault("메모", "")
        if row["메모"]:
            continue
        for keyword, memo in memo_rules:
            if keyword in row["적요"]:
                row["메모"] = memo
                break


def load_rules():
    rules = []
    with open(RULES_PATH, encoding="utf-8-sig") as f:
        for item in csv.DictReader(f):
            keyword = (item.get("키워드") or "").strip()
            if not keyword:
                continue
            rules.append((keyword, (item.get("대분류") or "").strip(), (item.get("소분류") or "").strip()))
    return rules


def categorize(rows, rules):
    for row in rows:
        if row.get("대분류"):  # 파서가 이미 확정한 카테고리는 건드리지 않음
            continue
        text = f"{row['적요']} {row['내용']} {row.get('메모', '')}".lower()
        for keyword, main, sub in rules:
            if keyword.lower() in text:
                row["대분류"], row["소분류"] = main, sub
                break
        else:
            if row["입금"]:
                row["대분류"], row["소분류"] = "수입", "미분류"
            else:
                row["대분류"], row["소분류"] = "미분류", "미분류"


def _date_str(value):
    """날짜값을 'YYYY-MM-DD' 문자열로 통일한다.
    (엑셀에 한 번 저장했다가 다시 읽으면 date가 datetime으로 바뀌는 등
    타입이 섞이기 때문에, 비교 전에 반드시 문자열로 맞춰줘야 중복 판정이 정확해진다.)"""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def _time_str(value):
    """시각값을 'HH:MM' 문자열로 통일한다."""
    if isinstance(value, datetime.datetime):
        value = value.time()
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.timedelta):
        minutes = int(value.total_seconds() // 60)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    return str(value)


def _amount(value):
    if value in (None, ""):
        return 0
    return int(value)


def make_key(row):
    return (
        _date_str(row["거래일자"]),
        _time_str(row["거래시간"]),
        row["적요"],
        _amount(row["출금"]),
        _amount(row["입금"]),
    )


def load_existing_rows():
    if not OUTPUT_PATH.exists():
        return []
    wb = openpyxl.load_workbook(OUTPUT_PATH)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or all(v is None for v in raw):
            continue
        rows.append(dict(zip(header, raw)))
    return rows


def _combine_datetime(date_value, time_value):
    """거래일자/거래시간을 하나의 datetime으로 합친다.
    (B열, C열 둘 다 이 값을 그대로 넣고, 표시 형식만 날짜/시간으로 다르게 준다)"""
    if isinstance(date_value, datetime.datetime):
        d = date_value.date()
    elif isinstance(date_value, datetime.date):
        d = date_value
    else:
        d = datetime.datetime.strptime(str(date_value)[:10], "%Y-%m-%d").date()

    if isinstance(time_value, datetime.datetime):
        t = time_value.time()
    elif isinstance(time_value, datetime.time):
        t = time_value
    elif isinstance(time_value, datetime.timedelta):
        t = (datetime.datetime.min + time_value).time()
    else:
        t = datetime.time(0, 0)

    return datetime.datetime.combine(d, t)


def write_output(rows):
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "가계부"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        combined = _combine_datetime(row["거래일자"], row["거래시간"])
        values = []
        for col in COLUMNS:
            if col in ("거래일자", "거래시간"):
                values.append(combined)
            else:
                values.append(row.get(col))
        ws.append(values)

    date_col = COLUMNS.index("거래일자") + 1
    time_col = COLUMNS.index("거래시간") + 1
    out_col = COLUMNS.index("출금") + 1
    in_col = COLUMNS.index("입금") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=date_col).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=time_col).number_format = "hh:mm"
        ws.cell(row=r, column=out_col).number_format = "#,##0"
        ws.cell(row=r, column=in_col).number_format = "#,##0"

    widths = [10, 12, 9, 26, 12, 12, 32, 14, 14, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # 1행에 자동 필터, 2행을 기준으로 틀 고정(1행 고정)
    last_col = openpyxl.utils.get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_PATH)


def main():
    print("=== 가계부 업데이트 시작 ===")
    print(f"입력 폴더: {INPUT_DIR}")

    new_rows, lookup_entries = discover_and_load()
    if not new_rows:
        print("새로 처리할 거래가 없습니다. 종료합니다.")
        return

    # 쿠팡/네이버페이/구글플레이 상품명, 고정 가맹점명 매핑은
    # AUTO_FILL 설정과 상관없이 항상 '메모' 칸에 채워준다.
    # 우선순위: 실제 주문내역으로 찾은 상품명(enrich_memo) > memo_mapping.csv
    # 고정 매핑(apply_memo_rules). 실제 주문내역을 먼저 적용해서, 있으면 그걸
    # 쓰고 그래도 비어있는 칸만 고정 매핑으로 채운다.
    amount_lookup = build_amount_lookup(lookup_entries)
    memo_rules = load_memo_rules()
    enrich_memo(new_rows, amount_lookup)
    apply_memo_rules(new_rows, memo_rules)

    if AUTO_FILL:
        rules = load_rules()
        categorize(new_rows, rules)
    else:
        for row in new_rows:
            row["내용"] = ""
            row["대분류"] = ""
            row["소분류"] = ""

    existing_rows = load_existing_rows()

    # 메모가 비어있는 거래는 (새 거래인지 예전 거래인지 상관없이) 최신 규칙/자료로
    # 채운다. 메모가 이미 있는 행은 절대 건드리지 않는다 (손으로 채운 메모 보호).
    # 우선순위는 위와 동일: 실제 주문내역(enrich_memo) > 고정 매핑(apply_memo_rules).
    backfilled_before = sum(1 for r in existing_rows if not r.get("메모"))
    enrich_memo(existing_rows, amount_lookup)
    apply_memo_rules(existing_rows, memo_rules)
    backfilled = backfilled_before - sum(1 for r in existing_rows if not r.get("메모"))

    existing_keys = {make_key(r) for r in existing_rows}

    added = [r for r in new_rows if make_key(r) not in existing_keys]
    skipped = len(new_rows) - len(added)

    all_rows = existing_rows + added
    # 1차: 월(연-월) 내림차순(최신 달이 위), 2차: 은행명 오름차순, 3차: 거래일시 내림차순
    # (정렬은 안정 정렬이므로, 낮은 우선순위부터 차례로 적용하면 마지막에 적용한
    #  기준이 최우선이 된다: 거래일시 -> 은행 -> 월 순서로 적용)
    all_rows.sort(key=lambda r: (_date_str(r["거래일자"]), _time_str(r["거래시간"])), reverse=True)
    all_rows.sort(key=lambda r: r["은행"])
    all_rows.sort(key=lambda r: _date_str(r["거래일자"])[:7], reverse=True)

    write_output(all_rows)

    print()
    print(f"이번에 읽은 거래: {len(new_rows)}건")
    print(f"이미 있어서 건너뜀(중복): {skipped}건")
    print(f"새로 추가됨: {len(added)}건")
    if backfilled:
        print(f"메모가 비어있던 과거 거래 중 새로 채운 건수: {backfilled}건")
    print(f"저장 위치: {OUTPUT_PATH}")

    if not AUTO_FILL:
        print()
        print("자동 분류가 꺼져있어 '내용/대분류/소분류'는 전부 빈칸으로 저장했습니다. "
              "가계부.xlsx에서 직접 채워주세요.")
    else:
        uncategorized = [r for r in added if r.get("대분류") in ("미분류", None)]
        if uncategorized:
            print()
            print(f"[확인 필요] 미분류로 남은 항목 {len(uncategorized)}건 (가계부.xlsx에서 직접 채워주세요):")
            for r in uncategorized[:30]:
                amt = r["출금"] or r["입금"]
                print(f"  - {r['거래일자']} {r['적요']} ({amt:,}원)")
        if len(uncategorized) > 30:
            print(f"  ... 외 {len(uncategorized) - 30}건 더 있음")
    print("=== 완료 ===")


if __name__ == "__main__":
    main()
